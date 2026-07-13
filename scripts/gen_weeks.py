"""
Parse the client's Overview workbook into structured per-week content for the
27 non-milestone weeks (5,6,7,9,10,11,...). The 10 milestone weeks (4,8,...,40)
are kept from the hand-authored src/weeks.js and are NOT regenerated here.

Each "Text Box" cell (EN + DE) has the shape:

    Baby
    <development paragraph>
    Mama
    <mama sentence>
    <emoji> <Label>: <callout text>

Output: scripts/weeks_parsed.json  (all 37 parsed, for inspection + the Node
merge step in gen_weeks_merge.js). Anomalies are printed so nothing is silently
mis-parsed.

Requires: openpyxl.  Run:  python scripts/gen_weeks.py
"""
import openpyxl, io, os, re, json

XLSX = r"C:\Users\jawad\Downloads\Pregnancy Calendar Overview.xlsx"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "scripts", "weeks_parsed.json")


def tokens(text):
    # Split on any run of newlines → clean tokens. Handles both the blank-line
    # (\n\n) blobs and the single-\n blobs (e.g. weeks 21, 26).
    return [t.strip() for t in re.split(r"\n+", (text or "").strip()) if t.strip()]


def is_emoji_led(p):
    return bool(p) and ord(p[0]) > 0x2000  # callout lines start with an emoji


def split_label(callout):
    """Split '<Label>: text' or '<Label>? text' → (label, text)."""
    c = re.sub(r"^\W+", "", callout).strip()   # strip leading emoji + spaces
    ci = c.find(":")
    qi = c.find("?")
    if ci != -1 and (qi == -1 or ci < qi):
        return c[:ci].strip(), c[ci + 1:].strip()
    if qi != -1:
        return c[:qi + 1].strip(), c[qi + 1:].strip()   # keep the '?' on the label
    return "", c


def parse_blob(text, week, lang, warns):
    ts = tokens(text)
    low = [t.lower() for t in ts]
    try:
        ib = low.index("baby")
        im = low.index("mama")
    except ValueError:
        warns.append(f"week {week} [{lang}]: missing Baby/Mama marker")
        return None
    dev = " ".join(ts[ib + 1:im]).strip()
    # After 'Mama': tokens up to the first emoji-led line are the Mama sentence;
    # each emoji-led line starts a new callout; a following non-emoji line (e.g.
    # 'Answer: ...') is a continuation of the current callout.
    mama_parts, callouts = [], []
    for t in ts[im + 1:]:
        if is_emoji_led(t):
            callouts.append(t)
        elif callouts:
            callouts[-1] += " " + t
        else:
            mama_parts.append(t)
    mama = " ".join(mama_parts).strip()
    if not callouts:
        warns.append(f"week {week} [{lang}]: no callout line")
        return {"dev": dev, "mama": mama, "label": "", "text": "", "tip": ""}
    label, ctext = split_label(callouts[0])
    if not label:
        warns.append(f"week {week} [{lang}]: callout label unparsed: {callouts[0][:40]!r}")
    # a 2nd callout line (Reminder / Milestone / To-Do) becomes the tip-block text
    tip = split_label(callouts[1])[1] if len(callouts) > 1 else ""
    return {"dev": dev, "mama": mama, "label": label, "text": ctext, "tip": tip}


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    warns, rows = [], []
    for r in range(2, 39):  # rows 2..38 = weeks 4..40
        week = int(round(float(ws.cell(r, 1).value)))
        en = parse_blob(ws.cell(r, 10).value, week, "en", warns)
        de = parse_blob(ws.cell(r, 11).value, week, "de", warns)
        if not en or not de:
            continue
        entry = {
            "week": week,
            "fruit": {"en": (ws.cell(r, 2).value or "").strip(),
                      "de": (ws.cell(r, 3).value or "").strip()},
            "length": (str(ws.cell(r, 6).value or "")).strip(),
            "weight": (str(ws.cell(r, 7).value or "")).strip(),
            "babyDev": {"en": en["dev"], "de": de["dev"]},
            "mama": {"en": en["mama"], "de": de["mama"]},
            "callout": {"label": {"en": en["label"], "de": de["label"]},
                        "text": {"en": en["text"], "de": de["text"]}},
        }
        if en["tip"] or de["tip"]:
            entry["tip"] = {"en": en["tip"], "de": de["tip"]}
        rows.append(entry)
    io.open(OUT, "w", encoding="utf-8").write(json.dumps(rows, ensure_ascii=False, indent=2))
    print(f"parsed {len(rows)} weeks -> {OUT}")
    print(f"\n{len(warns)} anomalies:")
    for w in warns:
        print("  -", w)


if __name__ == "__main__":
    main()
